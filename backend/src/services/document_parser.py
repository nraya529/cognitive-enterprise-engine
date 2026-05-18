from __future__ import annotations

import csv
import io
from typing import BinaryIO

from openpyxl import load_workbook
from pypdf import PdfReader


def parse_pdf(stream: BinaryIO) -> str:
    reader = PdfReader(stream)
    return "\n".join(page.extract_text() or "" for page in reader.pages).strip()


def parse_xlsx(stream: BinaryIO) -> str:
    wb = load_workbook(stream, read_only=True, data_only=True)
    lines: list[str] = []
    for sheet in wb.worksheets:
        lines.append(f"# Sheet: {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) if c is not None else "" for c in row]
            if any(cells):
                lines.append("\t".join(cells))
    return "\n".join(lines)


def parse_csv(text: str) -> str:
    reader = csv.reader(io.StringIO(text))
    return "\n".join("\t".join(row) for row in reader)


def parse_payload(content: bytes | str, content_type: str) -> str:
    if isinstance(content, str):
        if "csv" in content_type:
            return parse_csv(content)
        return content
    ct = content_type.lower()
    if "pdf" in ct:
        return parse_pdf(io.BytesIO(content))
    if "spreadsheet" in ct or "excel" in ct or "xlsx" in ct:
        return parse_xlsx(io.BytesIO(content))
    if "csv" in ct:
        return parse_csv(content.decode("utf-8", errors="replace"))
    return content.decode("utf-8", errors="replace")
